from openpyxl import load_workbook

# File Excel
input_file = "naver_미래_유망_직업_2021_2026_after_merge.xlsx"
output_file = "output.xlsx"

# Đọc workbook
wb = load_workbook(input_file)
ws = wb.active

deleted_count = 0
updated_count = 0

# Duyệt từ dưới lên để tránh bị lệch dòng khi xóa
for row in range(ws.max_row, 1, -1):  # Bỏ qua dòng tiêu đề
    value = ws[f"J{row}"].value

    if value is None or str(value).strip() == "":
        ws.delete_rows(row)
        deleted_count += 1
        print(f"[DELETE] Đã xóa dòng {row}")
    else:
        ws[f"K{row}"] = "Naver Blog"
        updated_count += 1
        print(f"[UPDATE] Dòng {row}: Thêm 'Naver Blog' vào cột K")

# Lưu file
wb.save(output_file)

print("\n========== KẾT QUẢ ==========")
print(f"Tổng số dòng đã xóa   : {deleted_count}")
print(f"Tổng số dòng đã cập nhật: {updated_count}")
print(f"File đã lưu: {output_file}")